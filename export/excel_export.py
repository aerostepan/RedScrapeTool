"""Excel workbook export using openpyxl."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any


RAW_HEADERS = [
    "Fetched At",
    "Published At",
    "Subreddit",
    "Query",
    "Type",
    "Title",
    "Text",
    "URL",
    "Author",
    "Reddit Score",
    "Comment Count",
]

FILTERED_HEADERS = [
    "Source",
    "Subreddit",
    "Title",
    "URL",
    "User Segment",
    "Pain Points",
    "Fears",
    "Needs",
    "Objections",
    "Desired Features",
    "Disliked Competitor Features",
    "Competitors Mentioned",
    "Existing Workarounds",
    "Pricing Complaints",
    "Integration Gaps",
    "Manual Workflows",
    "Emotional Intensity",
    "Urgency",
    "Willingness To Pay Signal",
    "Commercial Potential",
    "Confidence",
    "Summary",
    "Possible Startup Idea",
]

CLUSTER_HEADERS = [
    "Cluster Name",
    "Category",
    "Frequency",
    "Subreddits",
    "Competitors",
    "Representative Quotes",
    "Average Urgency",
    "Average Emotional Intensity",
    "Average Commercial Potential",
    "Workaround Count",
    "Pricing Complaint Count",
    "Integration Gap Count",
    "Opportunity Score",
]

FEATURE_HEADERS = [
    "Feature",
    "Frequency",
    "Related Pain Point",
    "Mentioned Competitors",
    "Representative Quotes",
    "Opportunity Score",
]

COMPETITOR_HEADERS = [
    "Competitor",
    "Weakness / Disliked Feature",
    "Frequency",
    "Related User Segment",
    "Representative Quotes",
    "Opportunity Score",
]

OPPORTUNITY_HEADERS = [
    "Rank",
    "Opportunity",
    "Main Pain Point",
    "Target User Segment",
    "Evidence Count",
    "Relevant Subreddits",
    "Competitors Mentioned",
    "Desired Features",
    "Workarounds Found",
    "Pricing Signals",
    "Integration Gaps",
    "Opportunity Score",
    "Confidence",
    "Representative Evidence URLs",
]


def export_market_research_workbook(
    output_path: str | Path,
    raw_items: list[dict[str, Any]],
    processed_items: list[dict[str, Any]],
    clusters: list[dict[str, Any]],
    desired_features: list[dict[str, Any]],
    competitor_weaknesses: list[dict[str, Any]],
    top_opportunities: list[dict[str, Any]],
) -> Path:
    """Create the multi-sheet Excel workbook."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise RuntimeError("openpyxl is required for Excel export. Run pip install -r requirements.txt") from exc

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    sheets = {
        "Raw Reddit Data": build_raw_rows(raw_items),
        "Filtered Signals": build_filtered_rows(processed_items),
        "Pain Point Clusters": build_cluster_rows(clusters),
        "Desired Features": build_feature_rows(desired_features),
        "Competitor Weaknesses": build_competitor_rows(competitor_weaknesses),
        "Top Opportunities": build_opportunity_rows(top_opportunities),
    }

    header_map = {
        "Raw Reddit Data": RAW_HEADERS,
        "Filtered Signals": FILTERED_HEADERS,
        "Pain Point Clusters": CLUSTER_HEADERS,
        "Desired Features": FEATURE_HEADERS,
        "Competitor Weaknesses": COMPETITOR_HEADERS,
        "Top Opportunities": OPPORTUNITY_HEADERS,
    }

    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(header_map[sheet_name])
        for row in rows:
            worksheet.append(row)
        style_worksheet(worksheet, get_column_letter, Font, PatternFill, Alignment)

    workbook.save(output)
    return output


def build_raw_rows(raw_items: list[dict[str, Any]]) -> list[list[Any]]:
    return [
        [
            item.get("fetched_at"),
            item.get("published_at"),
            item.get("subreddit"),
            item.get("query"),
            item.get("item_type"),
            item.get("title"),
            item.get("text"),
            item.get("url"),
            item.get("author"),
            item.get("score"),
            item.get("num_comments"),
        ]
        for item in raw_items
    ]


def build_filtered_rows(processed_items: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for item in processed_items:
        extraction = item.get("extraction") or {}
        if not extraction.get("is_relevant"):
            continue
        rows.append(
            [
                item.get("source"),
                item.get("subreddit"),
                item.get("title"),
                item.get("url"),
                extraction.get("user_segment"),
                join_cell(extraction.get("pain_points")),
                join_cell(extraction.get("fears")),
                join_cell(extraction.get("needs")),
                join_cell(extraction.get("objections")),
                join_cell(extraction.get("desired_features")),
                join_cell(extraction.get("disliked_competitor_features")),
                join_cell(extraction.get("competitors_mentioned")),
                join_cell(extraction.get("existing_workarounds")),
                join_cell(extraction.get("pricing_complaints")),
                join_cell(extraction.get("integration_gaps")),
                join_cell(extraction.get("manual_workflows")),
                extraction.get("emotional_intensity"),
                extraction.get("urgency"),
                extraction.get("willingness_to_pay_signal"),
                extraction.get("commercial_potential"),
                extraction.get("confidence"),
                extraction.get("summary"),
                extraction.get("possible_startup_idea"),
            ]
        )
    return rows


def build_cluster_rows(clusters: list[dict[str, Any]]) -> list[list[Any]]:
    return [
        [
            cluster.get("cluster_name"),
            cluster.get("category"),
            cluster.get("frequency"),
            join_cell(cluster.get("subreddits")),
            join_cell(cluster.get("competitors")),
            join_cell(cluster.get("representative_quotes")),
            cluster.get("avg_urgency"),
            cluster.get("avg_emotional_intensity"),
            cluster.get("avg_commercial_potential"),
            cluster.get("workaround_count"),
            cluster.get("pricing_complaint_count"),
            cluster.get("integration_gap_count"),
            cluster.get("opportunity_score"),
        ]
        for cluster in clusters
    ]


def build_feature_rows(features: list[dict[str, Any]]) -> list[list[Any]]:
    return [
        [
            item.get("feature"),
            item.get("frequency"),
            item.get("related_pain_point"),
            join_cell(item.get("mentioned_competitors")),
            join_cell(item.get("representative_quotes")),
            item.get("opportunity_score"),
        ]
        for item in features
    ]


def build_competitor_rows(weaknesses: list[dict[str, Any]]) -> list[list[Any]]:
    return [
        [
            item.get("competitor"),
            item.get("weakness"),
            item.get("frequency"),
            join_cell(item.get("related_user_segment")),
            join_cell(item.get("representative_quotes")),
            item.get("opportunity_score"),
        ]
        for item in weaknesses
    ]


def build_opportunity_rows(opportunities: list[dict[str, Any]]) -> list[list[Any]]:
    sorted_rows = sorted(opportunities, key=lambda item: item.get("opportunity_score") or 0, reverse=True)
    rows: list[list[Any]] = []
    for index, item in enumerate(sorted_rows, start=1):
        rows.append(
            [
                index,
                item.get("opportunity"),
                item.get("main_pain_point"),
                item.get("target_user_segment"),
                item.get("evidence_count"),
                join_cell(item.get("relevant_subreddits")),
                join_cell(item.get("competitors_mentioned")),
                join_cell(item.get("desired_features")),
                join_cell(item.get("workarounds_found")),
                join_cell(item.get("pricing_signals")),
                join_cell(item.get("integration_gaps")),
                item.get("opportunity_score"),
                item.get("confidence"),
                join_cell(item.get("representative_evidence_urls")),
            ]
        )
    return rows


def join_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(item) for item in value if str(item).strip())
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def style_worksheet(worksheet: Any, get_column_letter: Any, Font: Any, PatternFill: Any, Alignment: Any) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="000000")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_top

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_top

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value or "")
        max_length = len(header)
        for cell in column_cells[1:100]:
            value = str(cell.value or "")
            longest_line = max(value.split("\n"), key=len, default="")
            max_length = max(max_length, len(longest_line))
        if header in {"Text", "Summary", "Possible Startup Idea", "Representative Quotes"}:
            width = min(max(max_length + 2, 24), 70)
        elif "URL" in header:
            width = min(max(max_length + 2, 24), 60)
        else:
            width = min(max(max_length + 2, 12), 45)
        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 24

